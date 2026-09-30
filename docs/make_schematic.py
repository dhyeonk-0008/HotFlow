"""
HotFlow schematic (earlier alternative to docs/architecture.png).
Correct data flow: Receptor Features → GAEncoderCrossAttn (not HotspotEncoder).
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


# ── Helpers ──

def rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_box(left, top, w, h, fill, text='', fs=12, fc='#FFFFFF',
            bold=True, border=None, bw=Pt(1.5), align=PP_ALIGN.CENTER,
            sub=None, sub_fs=9, sub_fc=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(left), Inches(top), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if border:
        s.line.color.rgb = rgb(border)
        s.line.width = bw
    else:
        s.line.fill.background()
    s.adjustments[0] = 0.1
    tf = s.text_frame
    tf.word_wrap = True
    if text:
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(fs)
        r.font.color.rgb = rgb(fc)
        r.font.bold = bold
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = align
        p2.space_before = Pt(1)
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(sub_fs)
        r2.font.color.rgb = rgb(sub_fc or fc)
        r2.font.italic = True
        r2.font.bold = False
    return s


def add_txt(left, top, w, h, text, fs=10, fc='#333333',
            bold=False, italic=False, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(fs)
    r.font.color.rgb = rgb(fc)
    r.font.bold = bold
    r.font.italic = italic
    return tb


def add_multiline(left, top, w, h, lines, align=PP_ALIGN.CENTER):
    """lines: list of (text, fontsize, color, bold, italic)"""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, fs, fc, bold, italic) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(1)
        p.space_after = Pt(1)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(fs)
        r.font.color.rgb = rgb(fc)
        r.font.bold = bold
        r.font.italic = italic


def add_arrow(x1, y1, x2, y2, color='#2C3E50', w=Pt(2)):
    from pptx.oxml.ns import qn
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = w
    ln = c.line._ln
    te = ln.makeelement(qn('a:tailEnd'), {})
    te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    ln.append(te)
    return c


def add_dashed_arrow(x1, y1, x2, y2, color='#9B59B6', w=Pt(2)):
    from pptx.oxml.ns import qn
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = w
    ln = c.line._ln
    d = ln.makeelement(qn('a:prstDash'), {}); d.set('val', 'dash'); ln.append(d)
    te = ln.makeelement(qn('a:tailEnd'), {})
    te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
    ln.append(te)
    return c


# ── Colors ──
C = dict(
    rec='#4A90D9', pep='#7B68EE', anc='#FF6B6B', enc='#FF8C42', enc_s='#E07530',
    ipa='#5CB85C', ipa_x='#27AE60', xattn='#E74C3C', flow='#9B59B6',
    out='#1ABC9C', pred='#3498DB', cfg='#C0392B', txt='#2C3E50',
    s1='#E8EDF5', s2='#FFF5E8', info='#F0F0F0', loss='#ECEFF1',
    encode='#34495E',
)

# ══════════════════════════════════════════════
# TITLE
# ══════════════════════════════════════════════
add_txt(0, 0.1, 16, 0.5, 'HotFlow Architecture', fs=28, fc=C['txt'], bold=True)
add_txt(0, 0.55, 16, 0.3,
        'Hotspot-Conditioned Peptide Binder Design via SE(3) Flow Matching',
        fs=13, fc='#888888', italic=True)

# ══════════════════════════════════════════════
# STAGE 1 (left, x: 0.3–5.0)
# ══════════════════════════════════════════════
add_box(0.3, 1.0, 4.6, 7.7, C['s1'], border='#BBBBBB')
add_txt(0.5, 1.05, 3.5, 0.3, 'Stage 1: Hotspot Prediction',
        fs=12, fc=C['txt'], bold=True, italic=True, align=PP_ALIGN.LEFT)

# Receptor Pocket
add_box(0.8, 1.5, 3.6, 0.55, C['rec'], 'Receptor Pocket', fs=13)
add_arrow(2.6, 2.05, 2.6, 2.35)

# PepHAR EBM
add_box(0.8, 2.35, 3.6, 0.7, C['pep'], 'PepHAR EBM', fs=14,
        sub='pretrained & frozen', sub_fs=9, sub_fc='#D0C8E8')
add_arrow(2.6, 3.05, 2.6, 3.35)
add_txt(1.4, 3.05, 2.5, 0.25, 'density sampling', fs=8, fc='#888888', italic=True)

# K Hotspot Anchors
add_box(0.6, 3.4, 4.0, 0.85, C['anc'], 'K Hotspot Anchors', fs=13,
        sub='Ca/C/N coords + AA type', sub_fs=9)
add_txt(0.5, 4.3, 4.2, 0.3, '(B, K, 3, 3) coords  ·  (B, K) types  ·  K = 5',
        fs=9, fc=C['anc'], bold=True)

# Anchor Source info
add_box(0.5, 4.75, 4.2, 2.5, C['info'], border='#CCCCCC', bw=Pt(1))
add_multiline(0.7, 4.8, 3.8, 2.3, [
    ('Anchor Source', 11, C['txt'], True, False),
    ('', 4, '#F0F0F0', False, False),
    ('Training:   GT contacts (< 4 A)', 9.5, '#555555', False, True),
    ('Inference:  PepHAR EBM prediction', 9.5, '#555555', False, True),
    ('', 4, '#F0F0F0', False, False),
    ('Classifier-Free Guidance', 10, C['cfg'], True, False),
    ('p_drop = 0.1 (training)', 9.5, C['cfg'], False, False),
    ('guidance_scale > 1.0 (inference)', 9.5, C['cfg'], False, False),
], align=PP_ALIGN.LEFT)

# ══════════════════════════════════════════════
# STAGE 2 (right, x: 5.2–15.7)
# ══════════════════════════════════════════════
add_box(5.2, 1.0, 10.5, 7.7, C['s2'], border='#BBBBBB')
add_txt(5.4, 1.05, 6.0, 0.3, 'Stage 2: Hotspot-Conditioned Flow Matching',
        fs=12, fc=C['txt'], bold=True, italic=True, align=PP_ALIGN.LEFT)

# ── Top: Input batch → encode() ──
add_box(5.6, 1.55, 3.2, 0.7, C['encode'], 'encode(batch)', fs=12,
        sub='receptor + peptide structure', sub_fs=8, sub_fc='#AABBCC')

# encode() outputs: node_embed + edge_embed
add_txt(5.6, 2.3, 3.2, 0.3, 'node_embed (B,L,128)', fs=8, fc='#555555', bold=True)
add_txt(5.6, 2.5, 3.2, 0.3, 'edge_embed (B,L,L,64)', fs=8, fc='#555555', bold=True)

# Noisy State
add_box(12.0, 1.55, 3.2, 0.7, C['flow'], 'Noisy State xt', fs=13,
        sub='rot · trans · torsion · seq', sub_fs=8.5, sub_fc='#D8C8F0')
add_txt(12.6, 1.3, 2.0, 0.22, 't ~ U[0.01, 1]', fs=9, fc=C['flow'], bold=True)

# ── Middle-left: HotspotEncoder ──
# Arrow from Anchors → HotspotEncoder
add_arrow(4.65, 3.82, 5.55, 3.5, C['anc'], Pt(2.5))

add_box(5.6, 2.95, 3.5, 1.55, C['enc'], border=C['enc_s'])
add_txt(5.6, 2.98, 3.5, 0.3, 'HotspotEncoder', fs=13, fc='#FFFFFF', bold=True)

# Sub-modules 2x2
sw, sh = 1.5, 0.32
add_box(5.8, 3.35, sw, sh, C['enc_s'], 'SE(3) Inv.', fs=7.5, bold=False)
add_box(7.45, 3.35, sw, sh, C['enc_s'], 'AA Embed', fs=7.5, bold=False)
add_box(5.8, 3.75, sw, sh, C['enc_s'], 'MLP Proj.', fs=7.5, bold=False)
add_box(7.45, 3.75, sw, sh, C['enc_s'], 'Self-Attn', fs=7.5, bold=False)

# receptor_center annotation (small)
add_txt(5.2, 4.2, 2.0, 0.2, '+ receptor_center (B,3)', fs=7, fc='#999999', italic=True,
        align=PP_ALIGN.LEFT)

# Output: hotspot context
add_txt(5.6, 4.55, 3.5, 0.25, 'hotspot context  (B, 5, 64)',
        fs=10, fc=C['enc'], bold=True)

# ── IPA Blocks area ──
add_txt(9.6, 4.4, 5.5, 0.3, 'GAEncoderCrossAttn', fs=13, fc=C['txt'], bold=True)

# Background for IPA area
add_box(5.4, 4.8, 9.6, 2.55, '#E8E5DC', border='#CCCCCC', bw=Pt(1))

# Cross-Attn blocks (3, 4, 5) — top row within IPA area
bw_ipa = 1.35
gap = 0.15
bx0 = 5.65
by_cross = 5.0
by_ipa = 5.65
bh_ipa = 0.6
bh_cross = 0.45

for i in range(6):
    bx = bx0 + i * (bw_ipa + gap)
    has_x = i >= 3
    col = C['ipa_x'] if has_x else C['ipa']
    alpha = 1.0 if has_x else 0.85
    add_box(bx, by_ipa, bw_ipa, bh_ipa, col, f'IPA Block {i}', fs=9)
    add_txt(bx, by_ipa + bh_ipa + 0.02, bw_ipa, 0.18,
            '+ Seq Tfmr', fs=6.5, fc='#777777', italic=True)

    if has_x:
        add_box(bx, by_cross, bw_ipa, bh_cross, C['xattn'], 'Cross-Attn', fs=8)

    # Forward arrows between IPA blocks
    if i < 5:
        nbx = bx0 + (i + 1) * (bw_ipa + gap)
        add_arrow(bx + bw_ipa, by_ipa + bh_ipa / 2,
                  nbx, by_ipa + bh_ipa / 2, '#999999', Pt(1.5))

# ── Arrows into GAEncoderCrossAttn ──

# encode() outputs → IPA Block 0 (node_embed, edge_embed)
add_arrow(7.2, 2.85, 7.2, 4.75, C['encode'], Pt(2))

# Noisy state → IPA area (from top-right)
add_arrow(13.6, 2.25, 13.6, 4.75, C['flow'], Pt(2))

# HotspotEncoder → Cross-Attn blocks 3,4,5
for i in [3, 4, 5]:
    bx = bx0 + i * (bw_ipa + gap) + bw_ipa / 2
    add_arrow(7.35, 4.8, bx, by_cross, C['xattn'], Pt(1.5))

# Q/K/V label
add_txt(7.8, 4.6, 4.5, 0.2, 'Q: peptide nodes     K, V: hotspot context',
        fs=8, fc=C['xattn'], italic=True, align=PP_ALIGN.LEFT)

# ── Output row ──
last_x = bx0 + 5 * (bw_ipa + gap) + bw_ipa / 2
add_arrow(last_x, by_ipa + bh_ipa + 0.2, last_x, 7.5, C['txt'], Pt(2))

# Predicted clean state
add_box(10.8, 7.5, 4.2, 0.55, C['pred'], 'Predicted Clean State x1', fs=11,
        sub='trans · rot · torsion · AA probs', sub_fs=8, sub_fc='#B0D0FF')

# Denoising loop (dashed arrow: prediction → back to noisy state)
add_dashed_arrow(15.1, 7.5, 15.1, 2.25, C['flow'], Pt(2))
add_txt(15.15, 4.6, 0.8, 0.35, '100\nEuler\nsteps', fs=8, fc=C['flow'], bold=True)

# → Generated peptide
add_arrow(12.9, 8.05, 12.9, 8.3, C['out'], Pt(2.5))
add_box(10.3, 8.3, 5.2, 0.45, C['out'],
        'Generated Peptide (structure + sequence)', fs=11)

# ── Training Losses ──
add_box(5.5, 7.4, 4.5, 1.35, C['loss'], border='#CCCCCC', bw=Pt(1))
add_multiline(5.7, 7.42, 4.1, 1.3, [
    ('Training Losses', 10, C['txt'], True, False),
    ('trans (0.5)  ·  rot (0.5)  ·  bb_atom (0.25)', 8, '#666666', False, False),
    ('seq (1.0)  ·  angle (1.0)  ·  torsion (0.5)', 8, '#666666', False, False),
    ('contact_loss (0.1)', 10, C['xattn'], True, False),
])
add_arrow(10.8, 7.75, 10.05, 7.75, '#AAAAAA', Pt(1.2))

# ══════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════
pptx_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'schematic.pptx')
prs.save(pptx_path)
print(f"Saved: {pptx_path}")
